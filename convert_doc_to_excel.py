from docx import Document
from openpyxl import Workbook
import sys

def copy_doc_to_excel(doc_file):
    # Open the Word document
    doc = Document(doc_file)

    # Create a new workbook
    wb = Workbook()
    ws = wb.active

    # Add headers to the Excel sheet
    ws.append(["NO", "Content"])

    # Counter for the line numbers
    line_number = 1

    # Iterate through each paragraph in the document
    for para in doc.paragraphs:
        # Get the text of the paragraph
        text = para.text.strip()

        # Add the line to the Excel sheet
        ws.append([line_number, text])
        
        # Increment the line number
        line_number += 1

    # Extract the base filename
    base_filename = doc_file.split(".docx")[0]

    # Save the Excel file
    excel_file = f"{base_filename}.xlsx"
    wb.save(excel_file)
    print(f"Content from '{doc_file}' has been copied to '{excel_file}'.")

if __name__ == "__main__":
    if len(sys.argv) != 2:
        print("Usage: python script.py input.docx")
    else:
        doc_file = sys.argv[1]
        copy_doc_to_excel(doc_file)
